"""Parse ACMI submission templates (.xlsx). Returns preview rows only — no persistence."""

from __future__ import annotations

from io import BytesIO
from typing import Any


def parse_acmi_template(content: bytes) -> dict[str, Any]:
    if not content:
        return {
            "layoutMatched": False,
            "message": "The file is empty. No composition rows were read.",
            "rows": [],
        }
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("openpyxl is required to parse ACMI templates.") from error

    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception:
        return {
            "layoutMatched": False,
            "message": "This file is not a readable Excel workbook.",
            "rows": [],
        }

    # No sample ACMI layout is checked in yet — do not guess column positions.
    _ = workbook.sheetnames
    return {
        "layoutMatched": False,
        "message": "Layout not recognized. Provide a sample ACMI template to enable automatic import.",
        "rows": [],
    }
