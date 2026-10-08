"""Extract text from packaged source files. Results are cached by file hash."""

from __future__ import annotations

import json
import logging
from pathlib import Path

logger = logging.getLogger("affine.extract")

MIN_USEFUL_CHARS = 80


def _cache_path(data_dir: Path, digest: str) -> Path:
    folder = data_dir / ".extract_cache"
    folder.mkdir(parents=True, exist_ok=True)
    return folder / f"{digest}.json"


def _write_cache(path: Path, payload: dict) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def _read_cache(path: Path) -> dict | None:
    if not path.is_file():
        return None
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None
    return payload if isinstance(payload, dict) else None


def _extract_pdf(path: Path) -> list[dict]:
    try:
        from pypdf import PdfReader
    except ImportError:
        return [{"page": None, "text": "", "error": "pypdf_not_installed"}]
    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        pages.append({"page": index, "text": text, "locator_kind": "page"})
    return pages


def _extract_docx(path: Path) -> list[dict]:
    try:
        from docx import Document
    except ImportError:
        return [{"page": None, "text": "", "error": "python_docx_not_installed"}]
    document = Document(str(path))
    paragraphs = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = (paragraph.text or "").strip()
        if text:
            paragraphs.append({"page": None, "paragraph": index, "text": text, "locator_kind": "paragraph"})
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append({
                    "page": None,
                    "table": table_index,
                    "row": row_index,
                    "text": " | ".join(cells),
                    "locator_kind": "table_row",
                })
    return paragraphs


def extract_document(document: dict, data_dir: Path) -> dict:
    digest = str(document.get("sha256") or "").strip() or "unknown"
    cache = _cache_path(data_dir, digest)
    cached = _read_cache(cache)
    if cached:
        return cached

    filename = Path(str(document.get("filename") or "")).name
    source_dir = (data_dir / "source_documents").resolve()
    path = (source_dir / filename).resolve()
    if source_dir not in path.parents or not path.is_file():
        payload = {
            "document_id": document.get("document_id"),
            "filename": filename,
            "status": "file_missing",
            "ocr_required": False,
            "parts": [],
        }
        _write_cache(cache, payload)
        return payload

    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            parts = _extract_pdf(path)
        elif suffix in {".docx", ".doc"}:
            parts = _extract_docx(path)
        else:
            parts = [{"page": None, "text": path.read_text(encoding="utf-8", errors="replace"), "locator_kind": "file"}]
    except Exception as error:
        logger.info("extraction_failed document_id=%s error=%s", document.get("document_id"), type(error).__name__)
        payload = {
            "document_id": document.get("document_id"),
            "filename": filename,
            "status": "extraction_failed",
            "ocr_required": suffix == ".pdf",
            "parts": [],
        }
        _write_cache(cache, payload)
        return payload

    combined = " ".join(str(part.get("text") or "") for part in parts).strip()
    ocr_required = suffix == ".pdf" and len(combined) < MIN_USEFUL_CHARS
    status = "extracted" if combined else ("ocr_required" if ocr_required else "empty")
    payload = {
        "document_id": document.get("document_id"),
        "filename": filename,
        "status": status,
        "ocr_required": ocr_required,
        "parts": parts,
    }
    _write_cache(cache, payload)
    return payload
